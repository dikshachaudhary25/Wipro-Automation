import os
import re
import shutil
import zipfile
from datetime import datetime, timedelta

import openpyxl

# =====================================================
# CONFIGURATION
# =====================================================

ARCHIVE_PATH = r"\\molina.mhc\shares\Diplomat\Operations\BatchResults\all\M360DI\Archive"

ONEDRIVE_FOLDER = os.path.join(
    os.path.expanduser("~"),
    "OneDrive - Molina Healthcare"
)

SHEET_NAME = "Extracts Outbound & Inbound"

START_ROW = 68
END_ROW = 124

# =====================================================
# PROCESSING DATE
# =====================================================

requested_date = os.environ.get("TRANSACTION_DATE")
if requested_date:
    try:
        processing_date = datetime.strptime(
            requested_date,
            "%Y-%m-%d"
        ).strftime("%Y%m%d")
    except ValueError as exc:
        raise ValueError(
            "TRANSACTION_DATE must use YYYY-MM-DD format."
        ) from exc
else:
    processing_date = (
        datetime.now() - timedelta(days=1)
    ).strftime("%Y%m%d")

workbook_date = datetime.strptime(
    processing_date,
    "%Y%m%d"
).strftime("%m_%d_%Y")

workbook_override = os.environ.get("TRANSACTION_WORKBOOK")
EXCEL_FILE = os.path.abspath(workbook_override) if workbook_override else os.path.join(
    ONEDRIVE_FOLDER,
    f"Autosys Daily Status_{workbook_date}.xlsx"
)
output_sheet_override = os.environ.get("TRANSACTION_WORKSHEET")

TEMP_FOLDER = os.path.join(
    os.environ.get("TEMP", os.path.expanduser("~")),
    "TransactionAutomation",
    "temp"
)

print(f"\nPROCESSING DATE: {processing_date}")

# =====================================================
# HELPERS
# =====================================================

def should_ignore_file(file_name):

    file_name = file_name.lower()

    ignore_patterns = [
        ".pdf.",
        "report",
        "summary"
    ]

    return any(
        pattern in file_name
        for pattern in ignore_patterns
    )


def count_transactions(file_path, file_name):

    try:

        if should_ignore_file(file_name):
            return 0

        with open(
            file_path,
            "r",
            errors="ignore"
        ) as f:

            lines = sum(1 for _ in f)

        if lines <= 3:
            return 1

        return max(lines - 1, 0)

    except Exception:
        return 0


def count_written_lines(file_path, file_name):

    try:

        if should_ignore_file(file_name):
            return 0

        with open(
            file_path,
            "r",
            errors="ignore"
        ) as f:

            return sum(1 for _ in f)

    except Exception:
        return 0


def count_with_header(file_path, file_name):

    try:

        if should_ignore_file(file_name):
            return 0

        with open(
            file_path,
            "r",
            errors="ignore"
        ) as f:

            lines = [line.strip() for line in f if line.strip()]

        if not lines:
            return 0

        return max(len(lines) - 1, 0)

    except Exception:
        return 0


def is_gibberish_file(contents):

    if not contents:
        return False

    if b"\x00" in contents:
        return True

    try:
        text = contents.decode("utf-8")
    except UnicodeDecodeError:
        text = contents.decode("cp1252", errors="replace")

    control_characters = sum(
        1
        for character in text
        if ord(character) < 32 and character not in "\t\n\r\f"
    )

    return (
        "\ufffd" in text
        or control_characters / len(text) > 0.05
    )


def count_new_range_transactions(row, job_name, processing_date):

    zip_matches = get_zip_files(job_name)

    if not zip_matches:
        return 0

    if row == 3 and job_name == "elig_d_d_ftp_ib_all_mmp_m360_agent_exception":
        zip_path = os.path.join(ARCHIVE_PATH, zip_matches[-1])

        with zipfile.ZipFile(zip_path, "r") as zip_ref:
            dated_entries = [
                entry
                for entry in zip_ref.infolist()
                if not entry.is_dir()
                and processing_date in os.path.basename(entry.filename).split(
                    ".dest.",
                    1
                )[0]
            ]

            if len(dated_entries) != 1:
                raise ValueError(
                    f"Expected one {processing_date} file in {zip_matches[-1]}, "
                    f"found {len(dated_entries)}."
                )

            content = zip_ref.read(dated_entries[0]).decode(
                "utf-8",
                errors="replace"
            )

        loaded_counts = re.findall(
            r"NUMBER\s+OF\s+(?:\w+\s+)*RECORDS\s+LOADED\s*[:=]\s*([\d,]+)",
            content,
            flags=re.IGNORECASE
        )

        if not loaded_counts:
            raise ValueError(
                f"'Number of records loaded' was not found in "
                f"{dated_entries[0].filename}."
            )

        return int(loaded_counts[-1].replace(",", ""))

    lines_to_subtract = 3 if row in (5, 8) else 2
    total_transactions = 0

    for zip_name in zip_matches:
        zip_path = os.path.join(ARCHIVE_PATH, zip_name)

        with zipfile.ZipFile(zip_path, "r") as zip_ref:
            for entry in zip_ref.infolist():
                if entry.is_dir() or should_ignore_file(entry.filename):
                    continue

                contents = zip_ref.read(entry)
                if is_gibberish_file(contents):
                    total_transactions += 1
                    continue

                lines = len(
                    contents.decode("utf-8", errors="replace").splitlines()
                )
                total_transactions += max(lines - lines_to_subtract, 0)

    return total_transactions


def get_zip_files(job_name):

    matches = []

    for item in os.listdir(ARCHIVE_PATH):

        if (
            f".{job_name}." in item
            and f".{processing_date}." in item
            and item.endswith(".zip")
        ):

            matches.append(item)

    return sorted(matches)

# =====================================================
# OPEN WORKBOOK
# =====================================================

if not os.path.isfile(EXCEL_FILE):
    raise FileNotFoundError(
        f"Workbook for processing date {processing_date} was not found: "
        f"{EXCEL_FILE}"
    )

wb = openpyxl.load_workbook(EXCEL_FILE)

source_ws = wb[SHEET_NAME]
output_sheet_name = output_sheet_override or f"Transaction Count {processing_date}"

if output_sheet_name in wb.sheetnames:
    ws = wb[output_sheet_name]
elif output_sheet_override:
    raise ValueError(f"Worksheet not found: {output_sheet_name}")
else:
    source_sheet_title = source_ws.title
    temporary_source_title = "TempSource"
    while temporary_source_title in wb.sheetnames:
        temporary_source_title += "_"

    source_ws.title = temporary_source_title
    try:
        ws = wb.copy_worksheet(source_ws)
    finally:
        source_ws.title = source_sheet_title

    ws.title = output_sheet_name

# =====================================================
# FIND COLUMNS
# =====================================================

job_col = None
transaction_col = None

for row in range(1, 10):

    for col in range(1, ws.max_column + 1):

        value = ws.cell(
            row=row,
            column=col
        ).value

        if value is None:
            continue

        value = str(value)

        if value.strip() == "Job Name":
            job_col = col

        if "No. of Transactions for Yesterday" in value:
            transaction_col = col

if job_col is None:
    raise Exception("Job Name column not found")

if transaction_col is None:
    raise Exception("Transaction column not found")

ws.cell(
    row=1,
    column=transaction_col
).value = (
    "No. of Transactions for Yesterday, Date:  "
    f"{datetime.strptime(processing_date, '%Y%m%d').strftime('%m/%d/%Y')}"
)

# =====================================================
# PROCESS JOBS
# =====================================================

for row in range(2, 68):

    job_name = ws.cell(
        row=row,
        column=job_col
    ).value

    if not job_name:
        continue

    job_name = str(job_name).strip()

    print("\n" + "=" * 70)
    print(f"ROW {row}")
    print(f"JOB {job_name}")

    total_transactions = count_new_range_transactions(
        row,
        job_name,
        processing_date
    )

    ws.cell(
        row=row,
        column=transaction_col
    ).value = total_transactions

    print(f"TOTAL = {total_transactions}")


for row in range(START_ROW, END_ROW + 1):

    job_name = ws.cell(
        row=row,
        column=job_col
    ).value

    if not job_name:
        continue

    job_name = str(job_name).strip()

    print("\n" + "=" * 70)
    print(f"ROW {row}")
    print(f"JOB {job_name}")

    zip_matches = get_zip_files(job_name)
    print(f"MATCHING ZIP FILES = {len(zip_matches)}")

    total_transactions = None

    # --------------------------------------------------
    # FIXED VALUE RULES
    # --------------------------------------------------

    if job_name == "elig_d_d_ftp_ib_all_mmp_m360_member_exception":
        total_transactions = 0

    elif job_name == "elig_d_d_ftp_ob_all_mapd_m360_connecture_dental_oec":
        total_transactions = 0

    elif job_name == "elig_d_d_ftp_ib_all_mmp_m360_mbr_comments":
        total_transactions = 3

    elif job_name == "elig_d_d_ftp_ib_mmp_rejectinputfile":
        total_transactions = 1

    # --------------------------------------------------
    # ZIP CONTENT FILE COUNT RULES
    # --------------------------------------------------

    elif job_name in [
        "elig_d_d_ftp_ib_all_mmp_m360_834_inc",
        "elig_d_d_ftp_ib_all_mmp_m360_marxcob",
        "elig_d_d_move_ob_all_mmp_m360_marxcob"
    ]:

        total_transactions = 0

        for zip_name in zip_matches:

            zip_path = os.path.join(
                ARCHIVE_PATH,
                zip_name
            )

            try:

                with zipfile.ZipFile(
                    zip_path,
                    "r"
                ) as zip_ref:

                    file_count = sum(
                        1
                        for f in zip_ref.namelist()
                        if not f.endswith("/")
                    )

                    total_transactions += file_count

            except Exception:
                pass

    # --------------------------------------------------
    # ZIP CONTENT FILE COUNT RULES
    # --------------------------------------------------

    elif job_name in [
        "elig_d_d_ftp_ib_all_mmp_m360_lockbox04",
        "elig_d_d_ftp_ib_all_mmp_m360_dental_response"
    ]:

        total_transactions = 0

        for zip_name in zip_matches:

            zip_path = os.path.join(
                ARCHIVE_PATH,
                zip_name
            )

            try:

                with zipfile.ZipFile(
                    zip_path,
                    "r"
                ) as zip_ref:

                    file_count = sum(
                        1
                        for f in zip_ref.namelist()
                        if not f.endswith("/")
                    )

                    total_transactions += file_count

            except Exception:
                pass

    # --------------------------------------------------
    # NO-HEADER FILE COUNT RULES
    # --------------------------------------------------

    elif job_name in [
        "elig_d_d_ib_ftp_mapd_DTRRD",
        "oth_d_d_ib_ftp_infoxing_MARXCOB"
    ]:

        total_transactions = 0

        if len(zip_matches) == 0:

            total_transactions = 0

        else:

            if os.path.exists(TEMP_FOLDER):

                shutil.rmtree(
                    TEMP_FOLDER,
                    ignore_errors=True
                )

            os.makedirs(
                TEMP_FOLDER,
                exist_ok=True
            )

            for zip_name in zip_matches:

                zip_path = os.path.join(
                    ARCHIVE_PATH,
                    zip_name
                )

                extract_folder = os.path.join(
                    TEMP_FOLDER,
                    zip_name.replace(".zip", "")
                )

                os.makedirs(
                    extract_folder,
                    exist_ok=True
                )

                try:

                    with zipfile.ZipFile(
                        zip_path,
                        "r"
                    ) as zip_ref:

                        zip_ref.extractall(
                            extract_folder
                        )

                except Exception:
                    continue

                for root, dirs, files in os.walk(
                    extract_folder
                ):

                    for file in files:

                        file_path = os.path.join(
                            root,
                            file
                        )

                        total_transactions += count_written_lines(
                            file_path,
                            file
                        )

    # --------------------------------------------------
    # HEADER-AWARE FILE COUNT RULES
    # --------------------------------------------------

    elif job_name == "elig_d_d_ftp_ib_all_mmp_m360_prov_response":

        total_transactions = 0

        if len(zip_matches) == 0:

            total_transactions = 0

        else:

            if os.path.exists(TEMP_FOLDER):

                shutil.rmtree(
                    TEMP_FOLDER,
                    ignore_errors=True
                )

            os.makedirs(
                TEMP_FOLDER,
                exist_ok=True
            )

            for zip_name in zip_matches:

                zip_path = os.path.join(
                    ARCHIVE_PATH,
                    zip_name
                )

                extract_folder = os.path.join(
                    TEMP_FOLDER,
                    zip_name.replace(".zip", "")
                )

                os.makedirs(
                    extract_folder,
                    exist_ok=True
                )

                try:

                    with zipfile.ZipFile(
                        zip_path,
                        "r"
                    ) as zip_ref:

                        zip_ref.extractall(
                            extract_folder
                        )

                except Exception:
                    continue

                for root, dirs, files in os.walk(
                    extract_folder
                ):

                    for file in files:

                        file_path = os.path.join(
                            root,
                            file
                        )

                        total_transactions += count_with_header(
                            file_path,
                            file
                        )

    # --------------------------------------------------
    # NORMAL PROCESSING
    # --------------------------------------------------

    if total_transactions is None:

        if len(zip_matches) == 0:

            total_transactions = 0

        else:

            if job_name in [
                "elig_d_d_ftp_ib_all_mmp_m360_extract1",
                "elig_d_d_ftp_ib_all_mmp_m360_extract2",
                "elig_d_d_ftp_ib_all_mmp_m360_agent_exception"
            ]:

                zip_matches = [zip_matches[-1]]

            if os.path.exists(TEMP_FOLDER):

                shutil.rmtree(
                    TEMP_FOLDER,
                    ignore_errors=True
                )

            os.makedirs(
                TEMP_FOLDER,
                exist_ok=True
            )

            total_transactions = 0

            for zip_name in zip_matches:

                zip_path = os.path.join(
                    ARCHIVE_PATH,
                    zip_name
                )

                extract_folder = os.path.join(
                    TEMP_FOLDER,
                    zip_name.replace(".zip", "")
                )

                os.makedirs(
                    extract_folder,
                    exist_ok=True
                )

                try:

                    with zipfile.ZipFile(
                        zip_path,
                        "r"
                    ) as zip_ref:

                        zip_ref.extractall(
                            extract_folder
                        )

                except Exception:
                    continue

                for root, dirs, files in os.walk(
                    extract_folder
                ):

                    for file in files:

                        file_path = os.path.join(
                            root,
                            file
                        )

                        total_transactions += count_transactions(
                            file_path,
                            file
                        )

    # --------------------------------------------------
    # ADJUSTMENTS
    # --------------------------------------------------

    if job_name == "elig_d_d_ftp_ob_all_mmp_m360_lockbox05":
        total_transactions -= 1

    if job_name == "elig_d_d_ftp_ib_all_mmp_m360_compliance_tracker":
        total_transactions -= 1

    if total_transactions < 0:
        total_transactions = 0

    # --------------------------------------------------
    # WRITE
    # --------------------------------------------------

    ws.cell(
        row=row,
        column=transaction_col
    ).value = total_transactions

    print(f"TOTAL = {total_transactions}")

# =====================================================
# SAVE
# =====================================================

output_file = EXCEL_FILE if workbook_override else os.path.join(
    ONEDRIVE_FOLDER,
    f"Jobs_Updated_{processing_date}.xlsx"
)

wb.save(output_file)

print("\nDONE")
print(f"\nSAVED TO:\n{output_file}")